# -*- coding: utf-8 -*-
"""校验生成的 DOCX：字体是否全为宋体、结构是否完整、表格与标题是否规范。"""
import os
from docx import Document
from docx.oxml.ns import qn

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    '愈缝SeamHeal-项目简介文档.docx')
doc = Document(PATH)

bad = []
stats = {'runs': 0, 'body_paras': 0, 'tables': 0, 'headings': 0}


def check_runs(paras, where):
    for p in paras:
        for r in p.runs:
            if not r.text.strip():
                continue
            stats['runs'] += 1
            rPr = r._element.rPr
            fonts = {}
            if rPr is not None and rPr.rFonts is not None:
                for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                    fonts[a] = rPr.rFonts.get(qn(a))
            if r.font.name != '\u5b8b\u4f53' or any(v != '\u5b8b\u4f53' for v in fonts.values()):
                bad.append((where, r.text[:40], r.font.name, fonts))


check_runs(doc.paragraphs, 'body')
stats['body_paras'] = len(doc.paragraphs)

for hi, p in enumerate(doc.paragraphs):
    if p.style is not None and p.style.name.startswith('Heading'):
        stats['headings'] += 1

for ti, t in enumerate(doc.tables):
    stats['tables'] += 1
    for row in t.rows:
        for c in row.cells:
            check_runs(c.paragraphs, 'table%d' % ti)

# 节与页码
secs = doc.sections
print('节数：', len(secs))
for i, s in enumerate(secs):
    ftxt = ''.join(p.text for p in s.footer.paragraphs)
    has_fld = 'PAGE' in s.footer.paragraphs[0]._p.xml if s.footer.paragraphs else False
    print('  第%d节 页脚有PAGE域：%s' % (i + 1, has_fld))

# TOC 域
xml = doc.element.xml
print('TOC 域存在：', 'TOC \\o' in xml)
print('页码起始设置：', 'pgNumType' in xml)

print('\n结构统计：')
for k, v in stats.items():
    print('  %s = %s' % (k, v))

print('\n非宋体 run 数量：', len(bad))
for b in bad[:15]:
    print('  ', b)

# 样式字体抽查
print('\n样式字体抽查：')
for nm in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'Table Grid'):
    try:
        st = doc.styles[nm]
        rPr = st.element.rPr
        ea = rPr.rFonts.get(qn('w:eastAsia')) if (rPr is not None and rPr.rFonts is not None) else None
        print('  %-12s name=%-6s eastAsia=%s' % (nm, st.font.name, ea))
    except KeyError:
        print('  %-12s <不存在>' % nm)

# 表格列数一致性
print('\n表格规格：')
for i, t in enumerate(doc.tables):
    widths = set(len(r.cells) for r in t.rows)
    print('  表%d: %d行 x %s列' % (i + 1, len(t.rows), widths.pop() if len(widths) == 1 else widths))

print('\n总段落数：', len(doc.paragraphs))
