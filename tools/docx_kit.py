# -*- coding: utf-8 -*-
"""共用排版工具：全文档统一宋体，含标题、正文、表格、目录、页码等封装。

供 build_brief.py 与 build_bp.py 等生成脚本复用。
"""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG = '宋体'

# 待处理标记，由后续段落消费：
#   pb  下一个标题另起一页
#   gap 下一个段落增加表后留白
_PENDING = {'pb': False, 'gap': False}


def take_gap():
    if _PENDING['gap']:
        _PENDING['gap'] = False
        return 6
    return 0


def set_run_font(run, size=12, bold=False, name=SONG, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), name)


def force_song_everywhere(doc):
    """把所有内置样式的字体统一为宋体。"""
    for st in doc.styles:
        try:
            st.font.name = SONG
            rPr = st.element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                rFonts.set(qn(attr), SONG)
        except Exception:
            pass


def P(doc, text='', size=12, bold=False, align=None, indent=True,
      space_before=0, space_after=4, line=1.3, color=None, style=None):
    p = doc.add_paragraph(style=style)
    fmt = p.paragraph_format
    fmt.line_spacing = line
    fmt.space_before = Pt(space_before + take_gap())
    fmt.space_after = Pt(space_after)
    if indent:
        fmt.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def H(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12, 4: 12}
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, size=sizes.get(level, 12), bold=True, color=RGBColor(0, 0, 0))
    fmt = p.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(14 if level == 1 else (10 if level == 2 else 8))
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Pt(0)
    if _PENDING['pb']:
        fmt.page_break_before = True
        _PENDING['pb'] = False
    return p


def BULLET(doc, text, size=12, level=0):
    p = doc.add_paragraph(style='List Bullet')
    fmt = p.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(take_gap())
    fmt.space_after = Pt(3)
    fmt.left_indent = Pt(size * 2 + level * size)
    fmt.first_line_indent = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def TABLE(doc, headers, rows, widths=None, font_size=10.5, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(6 + take_gap())
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.line_spacing = 1.3
        cp.paragraph_format.first_line_indent = Pt(0)
        cp.paragraph_format.keep_with_next = True
        cr = cp.add_run(caption)
        set_run_font(cr, size=10.5, bold=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(htxt)
        set_run_font(r, size=font_size, bold=True)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:val'), 'clear')
        shade.set(qn('w:fill'), 'EFEFEF')
        hdr[i]._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Pt(0)
            r = p.add_run(str(val))
            set_run_font(r, size=font_size)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    # 表后不插入空段落（否则会被单独挤到下一页产生空白页），
    # 改为记入标记，由下一个段落通过 space_before 消费。
    _PENDING['gap'] = True
    return t


def CODE(doc, text, size=10):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.line_spacing = 1.15
    fmt.space_before = Pt(4 + take_gap())
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(size * 2)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Pt(0)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:color'), '888888')
    pbdr.append(bottom)
    pPr.append(pbdr)


def page_break(doc=None):
    _PENDING['pb'] = True


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run()
    set_run_font(r, size=10.5)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'PAGE \\* arabic \\* MERGEFORMAT'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    for e in (f1, it, f2):
        r._r.append(e)


def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run()
    set_run_font(r, size=12)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '【请在 Word 中右键此处，选择“更新域”以生成目录】'
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    for e in (f1, it, f2, t, f3):
        r._r.append(e)


def set_page_start(section, start=1):
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), str(start))
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        cols.addprevious(pgNumType)
    else:
        sectPr.append(pgNumType)


def enable_update_fields(doc):
    settings = doc.settings.element
    uf = settings.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields')
    if uf is None:
        uf = OxmlElement('w:updateFields')
        settings.append(uf)
    uf.set(qn('w:val'), 'true')


def setup_page(section):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
